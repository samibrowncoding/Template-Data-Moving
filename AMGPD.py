#filepath for IRR Template:
IRR_filepath = r"C:\Users\Intern\Documents\Intern\QuaterEnd report\2. LP reporting\Claret Fund III IRR Template Q2 2024 v2.xlsx"


#filepath for AM PD GP worksheet:
AM_PD_GP_filepath = r"C:\Users\Intern\Documents\Intern\QuaterEnd report\2. LP reporting\Testing\Originals\CEGCF III Q124 AM PD GP.xlsx"


#filepath to save output document to:
Output_filepath = r"C:\Users\Intern\Documents\Intern\QuaterEnd report\2. LP reporting\Testing\Outputs\ampdoutput-final.xlsx"


#No need to fiddle with anything below here
#One of the more fiddlier codes



#Make sure columns headers and sheet name match up














import pandas as pd
from openpyxl import load_workbook
import xlwings as xw
import numpy as np


#Code below is all about extracting tables from IRR template so we can prepare and use it. 
source1 = pd.read_excel(IRR_filepath, sheet_name='Realized Proceeds Table', skiprows=3, usecols="B:Q")
exclude_columns = ["MoC", "Gross IRR"]
columns_to_multiply = source1.columns.difference(exclude_columns)
source1[columns_to_multiply] = source1[columns_to_multiply].map(lambda x: x * 1000 if pd.api.types.is_numeric_dtype(type(x)) else x) 
source2 = pd.read_excel(IRR_filepath, sheet_name='Schedule of Investments', skiprows=4, usecols="A:R")
source2 = source2.map(lambda x: x * 1000 if pd.api.types.is_numeric_dtype(type(x)) else x) 

#Finding places to split up data
table_starts = source1[source1["Company"] == "Company"].index.tolist()
table_starts.append(len(source1))


#splitting up tables in "Realized Proceeds Table"
Unrealised_investments = source1.iloc[0:(table_starts[0] - 2)].reset_index(drop=True)
Partially_realised_investments = source1.iloc[table_starts[0]+2: table_starts[1] -2].reset_index(drop=True)
Fully_realised_proceeds = source1.iloc[table_starts[1]+2:table_starts[2] -5].reset_index(drop=True)

#Finding places to split up data
loan_indices = source2[source2["Investments"] == "Loan Securities"].index.tolist()
convertible_indices = source2[source2["Investments"] == "Convertible Loan Securities"].index.tolist()
equity_indices = source2[source2["Investments"] == "Equity Securities"].index.tolist()
warrants_indices = source2[source2["Investments"] == "Warrants"].index.tolist()
table_start = loan_indices + convertible_indices + equity_indices + warrants_indices

#splitting up tables in "Schedule of Investments"
Loan_securities = source2.iloc[table_start[0]: table_start[1] - 2]
Convertible_securities = source2.iloc[table_start[1]: table_start[2] -2]
Equity_Securities = source2.iloc[table_start[2]: table_start[3] -2]
Warrants = source2.iloc[table_start[3]: ] 

#minusing Equity and Warrant cost from Investment cost
new_Unrealised_investments = Unrealised_investments.copy()
for i, row in new_Unrealised_investments.iterrows():
    company_name = str(row['Company']).strip()
    investment_cost = row['Investment Cost']
    matching_equity = Equity_Securities[Equity_Securities['Investments'].str.strip() == company_name]
    if not matching_equity.empty and 'Cost*.2' in matching_equity.columns:
        cost_2_value_equity = matching_equity['Cost*.2'].values[0]
        investment_cost -= cost_2_value_equity

    matching_warrant = Warrants[Warrants['Investments'].str.strip() == company_name]
    if not matching_warrant.empty and 'Cost*.2' in matching_warrant.columns:
        cost_2_value_warrant = matching_warrant['Cost*.2'].values[0]
        investment_cost -= cost_2_value_warrant
    new_Unrealised_investments.at[i, 'Investment Cost'] = investment_cost 
#print(new_Unrealised_investments)

#Loading in AM PD GP sheet and sorting data 
workbook = load_workbook(AM_PD_GP_filepath, data_only=False)
sheet = workbook["Status"]
data = sheet.iter_rows(min_row=4, max_row=sheet.max_row, max_col=21, values_only=False)
columns = [cell.value for cell in sheet[3][0:21]]
am_pd_gp_df = pd.DataFrame([[cell.value for cell in row] for row in data], columns=columns)
#print(am_pd_gp_df[['Company Name', 'Nominal / Principal']])


#These are the functions that cause the change in value. Do not touch.
#If Columns headers change, can go find it in here and change if need be.
def update_AMPDGP_loan_data(am_pd_gp_df, new_Unrealised_investments):

    loan_filtered_df = am_pd_gp_df[am_pd_gp_df['Instrument Name'].str.contains("Loan", case=False, na=False)]
    for index, row in loan_filtered_df.iterrows():
        company_name = str(row["Company Name"]).strip()
        match = new_Unrealised_investments[new_Unrealised_investments['Company'].str.strip() == company_name]
        
        if not match.empty:
            am_pd_gp_df.at[index, 'Nominal / Principal'] = match['Investment Cost'].values[0]
            am_pd_gp_df.at[index, 'OID / Arrangement Fees'] = match['Fees'].values[0]
            am_pd_gp_df.at[index, 'Repayment of Principal / Realized costs of equity'] = match['Principal Repayments / Disposals'].values[0]
            am_pd_gp_df.at[index, 'Received Capitalized PIK interests / Capital Gain of equity'] = (match['Realised Gain/(Loss) (Equity / Warrants)'].values[0] + match['PIK Interest Accrued'].values[0])
            am_pd_gp_df.at[index, 'Received Cash interests / Income of equity'] = match['Interest Received'].values[0]
            am_pd_gp_df.at[index, 'Write up/down/off / Provisions / Unrealized Value'] = (match['Realised Gain/(Loss) (Loans)'].values[0] + match['FX Gain / (Loss) and Provisions (Loans)'].values[0])
            am_pd_gp_df.at[index, 'Accrued PIK interests'] = match['PIK Interest Accrued'].values[0]

        match_fmv_loan = Loan_securities[Loan_securities['Investments'].str.strip() == company_name]
        match_fmv_convertible = Convertible_securities[Convertible_securities['Investments'].str.strip() == company_name]
        match_fmv_pik = Loan_securities[Loan_securities['Investments'].str.strip() == f"{company_name} - PIK"]


        if not match_fmv_loan.empty:
            am_pd_gp_df.at[index, 'Valuation / FMV'] = match_fmv_loan['Fair Value.1'].values[0]

        if not match_fmv_pik.empty:
            am_pd_gp_df.at[index, 'Valuation / FMV'] += match_fmv_pik['Fair Value.1'].values[0]
            

        if not match_fmv_convertible.empty:
            am_pd_gp_df.at[index, 'Valuation / FMV'] += match_fmv_convertible['Fair Value.1'].values[0]

def update_AMPDGP_equity_data(am_pd_gp_df, Equity_Securities):
    equity_filtered_df = am_pd_gp_df[am_pd_gp_df['Instrument Name'].str.contains("Equity", case=False, na=False)]
    
    for index, row in equity_filtered_df.iterrows():
        company_name = str(row["Company Name"]).strip()
        
        match = Equity_Securities[Equity_Securities['Investments'].str.strip() == company_name]
        
        if not match.empty:
            am_pd_gp_df.at[index, 'Nominal / Principal'] = match['Cost*.2'].values[0]
            am_pd_gp_df.at[index, 'Write up/down/off / Provisions / Unrealized Value'] = match['Unnamed: 17'].values[0] 
            am_pd_gp_df.at[index, 'Valuation / FMV'] = match['Fair Value.1'].values[0]

def update_AMPDGP_warrant_data(am_pd_gp_df, Warrants):
    warrant_filtered_df = am_pd_gp_df[am_pd_gp_df['Instrument Name'].str.contains("Warrant", case=False, na=False)]
    
    for index, row in warrant_filtered_df.iterrows():
        company_name = str(row["Company Name"]).strip()
        
        match = Warrants[Warrants['Investments'].str.strip() == company_name]
        
        if not match.empty:
            am_pd_gp_df.at[index, 'Valuation / FMV'] = match['Fair Value.1'].values[0]
            am_pd_gp_df.at[index, 'Write up/down/off / Provisions / Unrealized Value'] = match['Unnamed: 17'].values[0]

#Implementing the function to update values.
update_AMPDGP_loan_data(am_pd_gp_df, new_Unrealised_investments)
update_AMPDGP_loan_data(am_pd_gp_df, Partially_realised_investments)
update_AMPDGP_loan_data(am_pd_gp_df, Fully_realised_proceeds)
update_AMPDGP_equity_data(am_pd_gp_df, Equity_Securities)
update_AMPDGP_warrant_data(am_pd_gp_df, Warrants)
#print(am_pd_gp_df[['Company Name', 'Nominal / Principal']])

#data slicing, making sure to save number of columns
if len(am_pd_gp_df.columns) > 21:
    am_pd_gp_df=am_pd_gp_df.iloc[:,:21]

#reading data back to worksheet 
for row_idx, row in am_pd_gp_df.iterrows():
    for col_idx, value in enumerate(row, start=1):
        sheet.cell(row=row_idx+4, column=col_idx, value=value)

workbook.save(Output_filepath)
print("All done yippee")

