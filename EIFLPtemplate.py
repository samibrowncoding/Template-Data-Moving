#filepath for IRR Template:
IRR_filepath = r"C:\Users\Intern\Documents\Intern\QuaterEnd report\2. LP reporting\Claret Fund III IRR Template Q2 2024 v2.xlsx"


#filepath for EIF Template:
EIF_filepath = r"C:\Users\Intern\Documents\Intern\QuaterEnd report\2. LP reporting\Testing\Originals\CEGCF III Q124 EIF Report.xlsx"


#filepath to save output document to:
Output_filepath = r"C:\Users\Intern\Documents\Intern\QuaterEnd report\2. LP reporting\Testing\Outputs\EIFoutput-final.xlsx"






#No need to mess around below here

#Make sure columns headers and sheet name match up











import pandas as pd
from openpyxl import load_workbook
import xlwings as xw

#Reading in data and multiplying by 1000
source = pd.read_excel(IRR_filepath, sheet_name='Realized Proceeds Table', skiprows=3, usecols="B:Q")
exclude_columns = ["MoC", "Gross IRR"]
columns_to_multiply = source.columns.difference(exclude_columns)
source[columns_to_multiply] = source[columns_to_multiply].map(lambda x: x * 1000 if pd.api.types.is_numeric_dtype(type(x)) else x)

#Finding places to split up data
table_starts = source[source["Company"] == "Company"].index.tolist()
table_starts.append(len(source))

#Splitting up Realized Proceeds data
Unrealised_investments = source.iloc[0:(table_starts[0] - 2)].reset_index(drop=True)
Partially_realised_investments = source.iloc[table_starts[0]+2: table_starts[1] -2].reset_index(drop=True)
Fully_realised_proceeds = source.iloc[table_starts[1]+2:table_starts[2] -5].reset_index(drop=True)

#Reading in EIF template and sorting
workbook = load_workbook(EIF_filepath, data_only=False)
sheet = workbook["2. Portfolio summary"]
data = sheet.iter_rows(min_row=5, max_row=sheet.max_row, min_col=2, max_col=14, values_only=False)
columns = [cell.value for cell in sheet[4][1:14]]
eif_df = pd.DataFrame([[cell.value for cell in row] for row in data], columns=columns)

#Function to change all Values. Fairly simple one
def update_eif_data(eif_df, Unrealised_investments):
    for index,row in eif_df.iterrows():

        company_name = str(row["Company name"]).strip()

        match = Unrealised_investments[Unrealised_investments['Company'].str.strip()== company_name]
        if not match.empty:
            eif_df.at[index, 'Current Cost'] = match['Fair Market Value'].values[0] - match['Unrealised \nGain / (Loss) (Equity / Warrants)'].values[0]
            eif_df.at[index, 'Realised cost (partial realisation)'] = match['Principal Repayments / Disposals'].values[0]
            eif_df.at[index, 'Total original cost'] = match['Investment Cost'].values[0]
            eif_df.at[index, 'Total cash proceeds (incl. return of capital, profit and income)'] = match['Total Realised Value'].values[0]
            eif_df.at[index, 'Fair value '] = match['Fair Market Value'].values[0]


#Using function here on from all 3 tables
update_eif_data(eif_df, Unrealised_investments)
update_eif_data(eif_df, Partially_realised_investments)
update_eif_data(eif_df, Fully_realised_proceeds)

#Slicing data to keep right format
if len(eif_df.columns) > 13:
    eif_df=eif_df.iloc[:,:13]

#Reading Data back to worksheet
for row_idx, row in eif_df.iterrows():
    for col_idx, value in enumerate(row, start=2):
        sheet.cell(row=row_idx+5, column=col_idx, value=value)

workbook.save(Output_filepath)
print("All done yippee")


















