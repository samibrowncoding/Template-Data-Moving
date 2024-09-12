#filepath for IRR Template:
IRR_filepath = r"C:\Users\Intern\Documents\Intern\QuaterEnd report\2. LP reporting\Claret Fund III IRR Template Q2 2024 v2.xlsx"


#filepath for Company list to split up Provison loss/FX reevaluation:
Provision_loss_filepath = r"C:\Users\Intern\Documents\Intern\QuaterEnd report\2. LP reporting\ProvisionLoss.xlsx"


#filepath for updated AM PD GP worksheet to change Revenue, EBITDA, and Cash for BCI:
AM_PD_GP_filepath = r"C:\Users\Intern\Documents\Intern\QuaterEnd report\2. LP reporting\CEGCF III Q224 AM PD GP.xlsx"


#filepath for BCI Template:
BCI_filepath = r"C:\Users\Intern\Documents\Intern\QuaterEnd report\2. LP reporting\Testing\Originals\BCI Report - Q124 - CEGCF III.xlsx"


#filepath to save output document to:
Output_filepath = r"C:\Users\Intern\Documents\Intern\QuaterEnd report\2. LP reporting\Testing\Outputs\BCIoutput-final.xlsx"


#When adding a new Compnay, make sure to put in "Company Issuer" AND "Instrument"

#No need to go below here, most confusing code one

#Make sure columns headers and sheet name match up












import pandas as pd
from openpyxl import load_workbook
import xlwings as xw

#Reading in data
source1 = pd.read_excel(IRR_filepath, sheet_name='Realized Proceeds Table', skiprows=3, usecols="B:Q")
exclude_columns = ["MoC", "Gross IRR"]
columns_to_multiply = source1.columns.difference(exclude_columns)
source1[columns_to_multiply] = source1[columns_to_multiply].map(lambda x: x * 1000 if pd.api.types.is_numeric_dtype(type(x)) else x)
source2 = pd.read_excel(IRR_filepath, sheet_name='Schedule of Investments', skiprows=4, usecols="A:R")
source2 = source2.map(lambda x: x * 1000 if pd.api.types.is_numeric_dtype(type(x)) else x)

#Finding places to split Realized Proceeds Table data
table_starts = source1[source1["Company"] == "Company"].index.tolist()
table_starts.append(len(source1))

#Splitting up data
Unrealised_investments = source1.iloc[0:(table_starts[0] - 2)].reset_index(drop=True)
Partially_realised_investments = source1.iloc[table_starts[0]+2: table_starts[1] -2].reset_index(drop=True)
Fully_realised_proceeds = source1.iloc[table_starts[1]+2:table_starts[2] -5].reset_index(drop=True)

#Finding places to split up Schedule of Investments data
loan_indices = source2[source2["Investments"] == "Loan Securities"].index.tolist()
convertible_indices = source2[source2["Investments"] == "Convertible Loan Securities"].index.tolist()
equity_indices = source2[source2["Investments"] == "Equity Securities"].index.tolist()
warrants_indices = source2[source2["Investments"] == "Warrants"].index.tolist()
table_start = loan_indices + convertible_indices + equity_indices + warrants_indices

#Splitting up data
Loan_securities = source2.iloc[table_start[0]: table_start[1] - 2]
Convertible_securities = source2.iloc[table_start[1]: table_start[2] -2]
Equity_Securities = source2.iloc[table_start[2]: table_start[3] -2]
Warrants = source2.iloc[table_start[3]: ] 

#Reading in excel sheet to distinguish which companies have provision loss or FX reevaluation split
df = pd.read_excel(Provision_loss_filepath)
provision_loss_companies = df['Company'].tolist()

#Reading in AM PD GP data for Money data
companies_money = pd.read_excel(AM_PD_GP_filepath, sheet_name='Companies', skiprows=1, usecols="A:M")

#minusing equity and warrant from Investment cost
new_Unrealised_investments = Unrealised_investments.copy()
for i, row in new_Unrealised_investments.iterrows():
    company_name = str(row['Company']).strip()
    investment_cost = row['Investment Cost']
    matching_equity = Equity_Securities[Equity_Securities['Investments'].str.strip() == company_name]
    if not matching_equity.empty and 'Cost*.2' in matching_equity.columns:
        cost_2_value_equity = matching_equity['Cost*.2'].values[0]
        investment_cost -= cost_2_value_equity

    matching_warrant = Warrants[Warrants['Investments'].str.strip() == company_name]
    if not matching_warrant.empty and 'Cost*.2' in matching_warrant.columns:
        cost_2_value_warrant = matching_warrant['Cost*.2'].values[0]
        investment_cost -= cost_2_value_warrant

    new_Unrealised_investments.at[i, 'Investment Cost'] = investment_cost 

#reading in BCI template data
workbook = load_workbook(BCI_filepath, data_only=False)
sheet = workbook["BCI Report"]
data = sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=21, values_only=False)
columns = [cell.value for cell in sheet[1][0:21]]
bci_df = pd.DataFrame([[cell.value for cell in row] for row in data], columns=columns)

#Functions to match up names and columns and change data
def update_BCI_loan_data(bci_df, new_Unrealised_investments):

    loan_filtered_df = bci_df[bci_df['Instrument'].str.contains("Loan/Bond", case=False, na=False)]
    for index, row in loan_filtered_df.iterrows():
        company_name = str(row["Company (Issuer)"]).strip()
        match = new_Unrealised_investments[new_Unrealised_investments['Company'].str.strip() == company_name]
        if not match.empty:
            bci_df.at[index, 'Investment Cost'] = match['Investment Cost'].values[0]
            bci_df.at[index, 'Realised proceeds (Principal)'] = match['Principal Repayments / Disposals'].values[0]
            bci_df.at[index, 'Realised proceeds (Interest + Fees)'] = (match['Interest Received'].values[0] + match['Fees'].values[0])
            bci_df.at[index, 'Warrant/ Equity Gain'] = match['Realised Gain/(Loss) (Equity / Warrants)'].values[0]
            fx_gain_loss = match['FX Gain / (Loss) and Provisions (Loans)'].values[0]
            if company_name in provision_loss_companies:
                bci_df.at[index, 'Loan write downs/provisions'] = fx_gain_loss
            else:
                bci_df.at[index, 'FX Revaluation(s)'] = fx_gain_loss

        match_fmv_loan = Loan_securities[Loan_securities['Investments'].str.strip() == company_name]
        match_fmv_convertible = Convertible_securities[Convertible_securities['Investments'].str.strip() == company_name]
        match_fmv_pik = Loan_securities[Loan_securities['Investments'].str.strip() == f"{company_name} - PIK"]


        if not match_fmv_loan.empty:
            bci_df.at[index, 'Fair Market value'] = match_fmv_loan['Fair Value.1'].values[0]
        if not match_fmv_pik.empty:
            bci_df.at[index, 'Fair Market value'] += match_fmv_pik['Fair Value.1'].values[0]
        
        if not match_fmv_convertible.empty:
            bci_df.at[index, 'Fair Market value'] += match_fmv_convertible['Fair Value.1'].values[0]

        match_companies_money = companies_money[companies_money['Company Name'].str.strip() == company_name]
        if not match_companies_money.empty:
            bci_df.at[index, 'Revenue (EUR)'] = match_companies_money['Revenue / Sales'].values[0]
            bci_df.at[index, 'EBITDA (EUR)'] = match_companies_money['EBITDA'].values[0]
            bci_df.at[index, 'Cash (EUR)'] = match_companies_money['Cash'].values[0]

def update_BCI_equity_data(bci_df, Equity_Securities):
    equity_filtered_df = bci_df[bci_df['Instrument'].str.contains("Equity", case=False, na=False)]
    
    for index, row in equity_filtered_df.iterrows():
        company_name = str(row["Company (Issuer)"]).strip()
        
        match = Equity_Securities[Equity_Securities['Investments'].str.strip() == company_name]
        
        if not match.empty:
            bci_df.at[index, 'Investment Cost'] = match['Cost*.2'].values[0]
            bci_df.at[index, 'Fair Market value'] = match['Fair Value.1'].values[0] 

def update_BCI_warrant_data(bci_df, Warrants):
    equity_filtered_df = bci_df[bci_df['Instrument'].str.contains("Warrant", case=False, na=False)]
    
    for index, row in equity_filtered_df.iterrows():
        company_name = str(row["Company (Issuer)"]).strip()
        
        match = Warrants[Warrants['Investments'].str.strip() == company_name]
        
        if not match.empty:
            bci_df.at[index, 'Investment Cost'] = match['Cost*.2'].values[0]
            bci_df.at[index, 'Fair Market value'] = match['Fair Value.1'].values[0] 

#Using functions to update data
update_BCI_loan_data(bci_df, new_Unrealised_investments)
update_BCI_loan_data(bci_df, Partially_realised_investments)
update_BCI_loan_data(bci_df, Fully_realised_proceeds)
update_BCI_equity_data(bci_df, Equity_Securities)
update_BCI_warrant_data(bci_df, Warrants)

#Reading data back to worksheet
for row_idx, row in bci_df.iterrows():
    for col_idx, value in enumerate(row, start=1):
        sheet.cell(row=row_idx+2, column=col_idx, value=value)

workbook.save(Output_filepath)
print("All done yippee")